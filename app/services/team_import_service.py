"""
Team import from .xlsx, per event (admin only).

Flow: template -> preview (read-only) -> confirm (all-or-nothing transaction).
Preview and confirm share parse_workbook() and build_report(), so a file that
previews clean cannot be rejected by a different code path at confirm time.
"""
import hashlib
import hmac
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Optional

from sqlalchemy.ext.asyncio import AsyncSession

from app.enums.education_level import EDUCATION_LEVELS
from app.models.team import Team
from app.repositories import event as event_repo
from app.repositories import team as team_repo
from app.utils import new_id

# Header row is the single source of truth: used by the template, the parser and tests.
IMPORT_COLUMNS: tuple[str, ...] = ("nama_tim", "sekolah", "kategori", "deskripsi", "urutan")
REQUIRED_COLUMNS: tuple[str, ...] = ("nama_tim", "kategori")

MAX_FILE_BYTES = 1024 * 1024
MAX_DATA_ROWS = 500
# Physical rows (including blank ones) read before giving up; bounds work on sheets padded with blanks.
MAX_PHYSICAL_ROWS = 5000
MAX_NAME = 200
MAX_SCHOOL = 200
MAX_DESCRIPTION = 2000
MAX_URUTAN = 100000

XLSX_MAGIC = b"PK\x03\x04"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE_SHEET = "Peserta"

# Labels agreed for Excel and the admin UI. Kept here on purpose: EDUCATION_LABELS is
# still the older public label set and is aligned in a separate step.
IMPORT_CATEGORY_LABELS: dict[str, str] = {
    "sd-mi": "SD/MI",
    "smp-mts": "SMP/MTs",
    "sma-smk-ma": "SMK/SMA/MA",
    "purna-klub": "PURNA/KLUB",
}

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_INTEGER_TEXT = re.compile(r"\d{1,6}")


def _fold(text: str) -> str:
    """Trim, collapse inner whitespace, lowercase. Used for headers, labels and identity."""
    return " ".join(text.split()).lower()


# Accepts each slug or its agreed label (case-insensitive, spacing normalized).
_CATEGORY_LOOKUP: dict[str, str] = {slug: slug for slug in EDUCATION_LEVELS}
_CATEGORY_LOOKUP.update({_fold(label): slug for slug, label in IMPORT_CATEGORY_LABELS.items()})

# Matches POST /api/v1/admin/events/{event_id}/teams/import/{preview|confirm}.
IMPORT_PATH_RE = re.compile(r"^/api/v1/admin/events/[^/]+/teams/import/(preview|confirm)$")


def is_import_path(method: str, path: str) -> bool:
    return method == "POST" and IMPORT_PATH_RE.match(path) is not None


class ImportFileError(Exception):
    """File-level problem that stops parsing entirely (wrong format, bad header, too many rows)."""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


class ImportRejected(Exception):
    """Business rejection returned to the client with an HTTP status and optional data."""

    def __init__(self, status_code: int, message: str, data: Any = None) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.message = message
        self.data = data


@dataclass
class ParsedRow:
    row_number: int
    values: dict[str, Any]
    errors: list[dict[str, str]] = field(default_factory=list)


@dataclass
class ParsedFile:
    rows: list[ParsedRow]
    skipped_empty_rows: list[int]
    warnings: list[str]


def _err(code: str, message: str) -> dict[str, str]:
    return {"code": code, "message": message}


def sha256_hex(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _identity(name: str, school: Optional[str], category: Optional[str]) -> tuple[str, str, str]:
    return (_fold(name or ""), _fold(school or ""), (category or "").strip().lower())


# --- Cell normalization -----------------------------------------------------
def _clean_line(text: str) -> str:
    text = _CONTROL_CHARS.sub("", text.replace("\t", " ").replace("\r", " "))
    return " ".join(text.split())


def _clean_multiline(text: str) -> str:
    text = _CONTROL_CHARS.sub("", text.replace("\r\n", "\n").replace("\r", "\n"))
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _is_formula(raw: Any) -> bool:
    return isinstance(raw, str) and raw.lstrip().startswith("=")


def _text_value(raw: Any, column: str, multiline: bool = False) -> tuple[str, Optional[dict[str, str]]]:
    if _is_formula(raw):
        return "", _err("rumus_tidak_diizinkan", f"Kolom {column} berisi rumus. Isi nilai langsung, bukan rumus.")
    if raw is None:
        return "", None
    if isinstance(raw, bool):
        return "", _err("teks_tidak_valid", f"Isi kolom {column} tidak valid")
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    if isinstance(raw, (int, float)):
        text = str(raw)
    elif isinstance(raw, datetime):
        # openpyxl reads date-only cells as datetime; midnight keeps the date form.
        text = raw.date().isoformat() if (raw.hour, raw.minute, raw.second) == (0, 0, 0) else raw.isoformat()
    elif isinstance(raw, date):
        text = raw.isoformat()
    elif isinstance(raw, str):
        text = raw
    else:
        return "", _err("teks_tidak_valid", f"Isi kolom {column} tidak valid")
    return (_clean_multiline(text) if multiline else _clean_line(text)), None


def _urutan_value(raw: Any) -> tuple[int, Optional[dict[str, str]]]:
    bad = _err("urutan_tidak_valid", f"Urutan harus bilangan bulat 0 sampai {MAX_URUTAN}")
    if _is_formula(raw):
        return 0, _err("rumus_tidak_diizinkan", "Kolom urutan berisi rumus. Isi nilai langsung, bukan rumus.")
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return 0, None
    if isinstance(raw, bool):
        return 0, bad
    if isinstance(raw, float):
        if not raw.is_integer():
            return 0, bad
        raw = int(raw)
    if isinstance(raw, str):
        text = raw.strip()
        if not _INTEGER_TEXT.fullmatch(text):
            return 0, bad
        raw = int(text)
    if not isinstance(raw, int) or raw < 0 or raw > MAX_URUTAN:
        return 0, bad
    return raw, None


def _category_value(raw: Any) -> tuple[Optional[str], Optional[dict[str, str]]]:
    text, err = _text_value(raw, "kategori")
    if err:
        return None, err
    if not text:
        return None, _err("kategori_wajib", "Kategori wajib diisi")
    slug = _CATEGORY_LOOKUP.get(_fold(text))
    if slug is None:
        shown = text if len(text) <= 40 else text[:40] + "…"
        options = ", ".join(IMPORT_CATEGORY_LABELS[s] for s in EDUCATION_LEVELS)
        return None, _err(
            "kategori_tidak_dikenal",
            f"Kategori tidak dikenal: '{shown}'. Pilih salah satu: {options}",
        )
    return slug, None


# --- Parsing (pure, no database) ---------------------------------------------
def _parse_row(row_number: int, cells: tuple, col: dict[str, int]) -> ParsedRow:
    def cell(name: str) -> Any:
        index = col[name]
        return cells[index] if index < len(cells) else None

    errors: list[dict[str, str]] = []

    name, err = _text_value(cell("nama_tim"), "nama_tim")
    if err:
        errors.append(err)
    elif not name:
        errors.append(_err("nama_wajib", "Nama tim wajib diisi"))
    elif len(name) > MAX_NAME:
        errors.append(_err("nama_terlalu_panjang", f"Nama tim maksimal {MAX_NAME} karakter"))

    school, err = _text_value(cell("sekolah"), "sekolah")
    if err:
        errors.append(err)
    elif len(school) > MAX_SCHOOL:
        errors.append(_err("sekolah_terlalu_panjang", f"Sekolah maksimal {MAX_SCHOOL} karakter"))

    category, err = _category_value(cell("kategori"))
    if err:
        errors.append(err)

    description, err = _text_value(cell("deskripsi"), "deskripsi", multiline=True)
    if err:
        errors.append(err)
    elif len(description) > MAX_DESCRIPTION:
        errors.append(_err("deskripsi_terlalu_panjang", f"Deskripsi maksimal {MAX_DESCRIPTION} karakter"))

    urutan, err = _urutan_value(cell("urutan"))
    if err:
        errors.append(err)

    return ParsedRow(
        row_number=row_number,
        values={
            "nama_tim": name,
            "sekolah": school,
            "kategori": category,
            "deskripsi": description,
            "urutan": urutan,
        },
        errors=errors,
    )


def _is_blank(raw: Any) -> bool:
    return raw is None or (isinstance(raw, str) and not raw.strip())


def parse_workbook(content: bytes) -> ParsedFile:
    """Parse the first sheet of an .xlsx. Never evaluates formulas; formula cells are rejected."""
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError("Ukuran file melebihi 1 MiB")
    if not content.startswith(XLSX_MAGIC):
        raise ImportFileError("Hanya file .xlsx yang didukung (.xls dan .csv tidak diterima)")

    import openpyxl

    # data_only=False keeps formula cells as "=..." strings so _is_formula can
    # reject them; with data_only=True a cached formula result would be imported
    # silently, which the plan forbids.
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ImportFileError("File tidak dapat dibaca sebagai .xlsx") from exc

    try:
        if not workbook.worksheets:
            raise ImportFileError("File tidak memiliki sheet")
        rows_iter = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None or all(_is_blank(c) for c in header):
            raise ImportFileError("Baris header tidak ditemukan di baris pertama")

        col: dict[str, int] = {}
        extra: list[str] = []
        for index, raw in enumerate(header):
            name = _fold(str(raw)) if raw is not None else ""
            if not name:
                continue
            if name in IMPORT_COLUMNS:
                if name in col:
                    raise ImportFileError(f"Kolom '{name}' muncul lebih dari sekali di header")
                col[name] = index
            else:
                extra.append(name)

        missing = [name for name in IMPORT_COLUMNS if name not in col]
        if missing:
            raise ImportFileError(
                "Header tidak sesuai template. Kolom yang hilang: " + ", ".join(missing)
            )

        warnings = []
        if extra:
            warnings.append("Kolom tambahan diabaikan: " + ", ".join(extra))

        rows: list[ParsedRow] = []
        skipped: list[int] = []
        physical = 0
        for row_number, cells in enumerate(rows_iter, start=2):
            physical += 1
            if physical > MAX_PHYSICAL_ROWS:
                raise ImportFileError(f"Sheet terlalu panjang (maksimal {MAX_PHYSICAL_ROWS} baris)")
            if all(_is_blank(c) for c in cells):
                skipped.append(row_number)
                continue
            rows.append(_parse_row(row_number, cells, col))
            if len(rows) > MAX_DATA_ROWS:
                raise ImportFileError(f"Jumlah baris data melebihi {MAX_DATA_ROWS}")
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError("File .xlsx rusak atau tidak dapat dibaca") from exc
    finally:
        workbook.close()

    return ParsedFile(rows=rows, skipped_empty_rows=skipped, warnings=warnings)


# --- Evaluation (duplicates, event category) --------------------------------
def build_report(
    parsed: ParsedFile,
    event_category: Optional[str],
    existing: set[tuple[str, str, str]],
    sha256: str,
) -> dict[str, Any]:
    counts = {"valid": 0, "invalid": 0, "duplicate_in_file": 0, "duplicate_existing": 0}
    seen: dict[tuple[str, str, str], int] = {}
    out_rows = []

    for parsed_row in parsed.rows:
        errors = list(parsed_row.errors)
        values = parsed_row.values
        category = values["kategori"]
        if event_category and category and category != event_category:
            errors.append(_err(
                "kategori_tidak_sesuai",
                f"Kategori tidak sesuai dengan event (diharapkan: {IMPORT_CATEGORY_LABELS[event_category]})",
            ))

        if errors:
            status = "invalid"
            counts["invalid"] += 1
        else:
            key = _identity(values["nama_tim"], values["sekolah"], category)
            if key in existing:
                status = "duplicate"
                counts["duplicate_existing"] += 1
                errors.append(_err(
                    "duplikat_event",
                    "Sudah ada tim dengan nama, sekolah, dan kategori yang sama di event ini",
                ))
            elif key in seen:
                status = "duplicate"
                counts["duplicate_in_file"] += 1
                errors.append(_err("duplikat_file", f"Duplikat dengan baris {seen[key]} dalam file"))
            else:
                seen[key] = parsed_row.row_number
                status = "valid"
                counts["valid"] += 1

        out_rows.append({
            "row_number": parsed_row.row_number,
            "status": status,
            "errors": errors,
            "data": {
                "nama_tim": values["nama_tim"],
                "sekolah": values["sekolah"],
                "kategori": category,
                "kategori_label": IMPORT_CATEGORY_LABELS.get(category or "", ""),
                "deskripsi": values["deskripsi"],
                "urutan": values["urutan"],
            },
        })

    blocking = None
    if event_category is None:
        blocking = "Atur kategori event terlebih dahulu"
    elif counts["valid"] == 0:
        blocking = "Tidak ada baris valid untuk diimpor"

    can_confirm = (
        blocking is None
        and counts["invalid"] == 0
        and counts["duplicate_in_file"] == 0
        and counts["duplicate_existing"] == 0
    )
    return {
        "sha256": sha256,
        "event_category": event_category,
        "can_confirm": can_confirm,
        "blocking_message": blocking,
        "total_rows": len(parsed.rows),
        "valid_count": counts["valid"],
        "invalid_count": counts["invalid"],
        "duplicate_in_file_count": counts["duplicate_in_file"],
        "duplicate_existing_count": counts["duplicate_existing"],
        "skipped_empty_rows": parsed.skipped_empty_rows,
        "warnings": parsed.warnings,
        "rows": out_rows,
    }


async def load_existing_identities(session: AsyncSession, event_id: str) -> set[tuple[str, str, str]]:
    rows = await team_repo.list_identity_rows(session, event_id)
    return {_identity(name, school, level) for name, school, level in rows}


# --- Public operations -------------------------------------------------------
def build_template(event_category: Optional[str]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    sheet.append(list(IMPORT_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    example_category = IMPORT_CATEGORY_LABELS.get(event_category or "", IMPORT_CATEGORY_LABELS["sd-mi"])
    sheet.append(["CONTOH - hapus baris ini", "Nama sekolah", example_category, "Deskripsi singkat (opsional)", 0])
    for letter, width in zip("ABCDE", (32, 28, 16, 40, 10)):
        sheet.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def preview_import(session: AsyncSession, event_id: str, event_category: Optional[str], content: bytes) -> dict[str, Any]:
    """Read-only: parses and evaluates the file, never writes teams."""
    try:
        parsed = parse_workbook(content)
    except ImportFileError as exc:
        raise ImportRejected(422, exc.message) from exc
    existing = await load_existing_identities(session, event_id)
    return build_report(parsed, event_category, existing, sha256_hex(content))


async def confirm_import(
    session: AsyncSession, event_id: str, content: bytes, expected_sha256: str
) -> list[Team]:
    """All-or-nothing insert. The caller commits on success and rolls back on ImportRejected."""
    event = await event_repo.get_event(session, event_id)
    if event is None or event.deleted_at is not None:
        raise ImportRejected(404, "Event tidak ditemukan")
    if event.category is None:
        raise ImportRejected(422, "Atur kategori event terlebih dahulu")
    if len(content) > MAX_FILE_BYTES:
        raise ImportRejected(422, "Ukuran file melebihi 1 MiB")

    actual = sha256_hex(content)
    if not hmac.compare_digest(actual, expected_sha256.lower()):
        raise ImportRejected(
            409, "File tidak sama dengan pratinjau terakhir. Unggah ulang file dan periksa pratinjau."
        )

    try:
        parsed = parse_workbook(content)
    except ImportFileError as exc:
        raise ImportRejected(422, exc.message) from exc

    # Lock the event row so two confirms for the same event run one after another.
    # The category is re-read under the lock; the duplicate check runs after it.
    found, category = await event_repo.lock_event_category(session, event_id)
    if not found:
        await session.rollback()
        raise ImportRejected(404, "Event tidak ditemukan")
    if category is None:
        await session.rollback()
        raise ImportRejected(422, "Atur kategori event terlebih dahulu")

    existing = await load_existing_identities(session, event_id)
    report = build_report(parsed, category, existing, actual)
    if not report["can_confirm"]:
        await session.rollback()
        problems = [row for row in report["rows"] if row["status"] != "valid"]
        raise ImportRejected(
            422,
            report["blocking_message"]
            or f"Impor dibatalkan: {len(problems)} baris perlu diperbaiki. Tidak ada tim yang disimpan.",
            data={"rows": problems, "invalid_count": report["invalid_count"],
                  "duplicate_in_file_count": report["duplicate_in_file_count"],
                  "duplicate_existing_count": report["duplicate_existing_count"]},
        )

    teams = []
    for row in report["rows"]:
        data = row["data"]
        teams.append(Team(
            id=new_id("tm"),
            event_id=event_id,
            name=data["nama_tim"],
            school=data["sekolah"] or None,
            education_level=data["kategori"],
            description=data["deskripsi"] or None,
            logo_url=None,
            sort_order=data["urutan"],
        ))
    session.add_all(teams)
    await session.flush()
    return teams
