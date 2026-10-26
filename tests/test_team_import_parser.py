"""Unit tests for the pure Excel import parser/validator (no database)."""
import io

import openpyxl
import pytest

from app.services import team_import_service as svc


def _xlsx(rows, headers=None, sheet_data_first=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    if headers is not False:
        ws.append(list(headers if headers is not None else svc.IMPORT_COLUMNS))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _valid_row(name="Tim A", school="SMA 1", category="SD/MI", desc=None, urutan=0):
    return [name, school, category, desc, urutan]


def test_valid_rows_parse():
    parsed = svc.parse_workbook(_xlsx([_valid_row(), _valid_row("Tim B", None, "smp-mts", "desc", 5)]))
    assert len(parsed.rows) == 2
    assert parsed.rows[0].errors == []
    assert parsed.rows[1].values["kategori"] == "smp-mts"
    assert parsed.rows[1].values["urutan"] == 5


def test_wrong_header_rejected():
    with pytest.raises(svc.ImportFileError) as exc:
        svc.parse_workbook(_xlsx([_valid_row()], headers=["nama", "sekolah", "kategori", "desc", "urut"]))
    assert "nama_tim" in exc.value.message


def test_header_case_and_trim():
    parsed = svc.parse_workbook(_xlsx(
        [_valid_row()],
        headers=["  NAMA_TIM ", "Sekolah", "KATEGORI", "Deskripsi", "Urutan"],
    ))
    assert parsed.rows[0].errors == []


def test_missing_category_column_rejected():
    with pytest.raises(svc.ImportFileError):
        svc.parse_workbook(_xlsx([["Tim A", "SMA 1"]], headers=["nama_tim", "sekolah"]))


def test_duplicate_header_rejected():
    with pytest.raises(svc.ImportFileError):
        svc.parse_workbook(_xlsx(
            [_valid_row()],
            headers=["nama_tim", "sekolah", "kategori", "deskripsi", "urutan", "NAMA_TIM"],
        ))


def test_extra_column_reported_as_warning():
    parsed = svc.parse_workbook(_xlsx(
        [_valid_row(), "ekstra"],
        headers=list(svc.IMPORT_COLUMNS) + ["kolom_aneh"],
    ))
    assert any("kolom_aneh" in w for w in parsed.warnings)


def test_unknown_category_rejected():
    parsed = svc.parse_workbook(_xlsx([_valid_row(category="Kuliah")]))
    codes = [e["code"] for e in parsed.rows[0].errors]
    assert "kategori_tidak_dikenal" in codes


def test_category_accepts_slug_and_label():
    parsed = svc.parse_workbook(_xlsx([
        _valid_row("A", category="sd-mi"),
        _valid_row("B", category="SMK/SMA/MA"),
        _valid_row("C", category="  purna / klub ".replace(" / ", "/")),
    ]))
    assert parsed.rows[0].values["kategori"] == "sd-mi"
    assert parsed.rows[1].values["kategori"] == "sma-smk-ma"
    assert parsed.rows[2].values["kategori"] == "purna-klub"


def test_missing_name_and_blank_name():
    parsed = svc.parse_workbook(_xlsx([_valid_row(name=None), _valid_row(name="   ")]))
    assert "nama_wajib" in [e["code"] for e in parsed.rows[0].errors]
    assert "nama_wajib" in [e["code"] for e in parsed.rows[1].errors]


def test_formula_cell_rejected():
    parsed = svc.parse_workbook(_xlsx([_valid_row(name="=A1+B1", desc="=SUM(A1:A2)"), _valid_row()]))
    codes = {e["code"] for e in parsed.rows[0].errors}
    assert "rumus_tidak_diizinkan" in codes


def test_number_in_name_column_coerced_to_text():
    parsed = svc.parse_workbook(_xlsx([_valid_row(name=12345)]))
    assert parsed.rows[0].values["nama_tim"] == "12345"


def test_date_cell_in_school_becomes_iso_text():
    from datetime import date

    parsed = svc.parse_workbook(_xlsx([_valid_row(school=date(2026, 1, 2))]))
    assert parsed.rows[0].values["sekolah"] == "2026-01-02"


def test_float_urutan_rejected_and_zero_default():
    parsed = svc.parse_workbook(_xlsx([_valid_row(urutan=1.5), _valid_row(), _valid_row(urutan=None)]))
    assert "urutan_tidak_valid" in [e["code"] for e in parsed.rows[0].errors]
    assert parsed.rows[1].values["urutan"] == 0
    assert parsed.rows[2].values["urutan"] == 0


def test_out_of_range_urutan_rejected():
    parsed = svc.parse_workbook(_xlsx([_valid_row(urutan=100001), _valid_row(urutan="abc")]))
    assert "urutan_tidak_valid" in [e["code"] for e in parsed.rows[0].errors]
    assert "urutan_tidak_valid" in [e["code"] for e in parsed.rows[1].errors]


def test_too_long_values_rejected():
    parsed = svc.parse_workbook(_xlsx([
        _valid_row(name="x" * 201, school="y" * 201, desc="z" * 2001),
    ]))
    codes = {e["code"] for e in parsed.rows[0].errors}
    assert {"nama_terlalu_panjang", "sekolah_terlalu_panjang", "deskripsi_terlalu_panjang"} <= codes


def test_blank_rows_skipped_and_reported():
    parsed = svc.parse_workbook(_xlsx([_valid_row(), (None, "", None, None, None), _valid_row("Tim B")]))
    assert len(parsed.rows) == 2
    assert parsed.skipped_empty_rows == [3]


def test_row_limit_enforced():
    rows = [_valid_row(f"Tim {i}") for i in range(svc.MAX_DATA_ROWS + 1)]
    with pytest.raises(svc.ImportFileError) as exc:
        svc.parse_workbook(_xlsx(rows))
    assert "500" in exc.value.message


def test_not_xlsx_rejected():
    with pytest.raises(svc.ImportFileError) as exc:
        svc.parse_workbook(b"id,nama\n1,x")
    assert ".xlsx" in exc.value.message

    with pytest.raises(svc.ImportFileError):
        svc.parse_workbook(b"PK\x03\x04 not really a zip")


def test_file_size_limit():
    with pytest.raises(svc.ImportFileError) as exc:
        svc.parse_workbook(svc.XLSX_MAGIC + b"\x00" * (svc.MAX_FILE_BYTES + 1))
    assert "1 MiB" in exc.value.message


def test_report_duplicate_and_category_rules():
    parsed = svc.parse_workbook(_xlsx([
        _valid_row("Tim A", "SMA 1", "SD/MI"),
        _valid_row("tim  a ", " sma 1 ", "sd-mi"),  # same identity, folded
        _valid_row("Tim C", None, "SMP/MTs"),
    ]))
    report = svc.build_report(parsed, "sd-mi", existing=set(), sha256="x")
    assert report["rows"][0]["status"] == "valid"
    assert report["rows"][1]["status"] == "duplicate"
    assert report["rows"][2]["status"] == "invalid"
    assert report["valid_count"] == 1
    assert report["duplicate_in_file_count"] == 1
    assert report["invalid_count"] == 1
    assert report["can_confirm"] is False


def test_report_existing_duplicate_and_event_mismatch():
    parsed = svc.parse_workbook(_xlsx([
        _valid_row("Tim A", "SMA 1", "SD/MI"),
        _valid_row("Tim D", None, "sd-mi"),
    ]))
    report = svc.build_report(parsed, "sma-smk-ma", existing=set(), sha256="x")
    # event category mismatch marks rows invalid even before duplicate checks
    assert report["invalid_count"] == 2
    assert "kategori_tidak_sesuai" in [e["code"] for e in report["rows"][0]["errors"]]


def test_report_blocks_null_event_category():
    parsed = svc.parse_workbook(_xlsx([_valid_row()]))
    report = svc.build_report(parsed, None, existing=set(), sha256="x")
    assert report["can_confirm"] is False
    assert report["blocking_message"] == "Atur kategori event terlebih dahulu"


def test_report_all_rows_blocked_by_existing_duplicates():
    parsed = svc.parse_workbook(_xlsx([_valid_row()]))
    report = svc.build_report(parsed, "sd-mi", existing={("tim a", "sma 1", "sd-mi")}, sha256="x")
    assert report["can_confirm"] is False


def test_template_has_expected_headers():
    wb = openpyxl.load_workbook(io.BytesIO(svc.build_template("smp-mts")))
    ws = wb.active
    assert list(next(ws.iter_rows(values_only=True))) == list(svc.IMPORT_COLUMNS)
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][0].startswith("CONTOH")
    assert rows[1][2] == "SMP/MTs"


def test_sha256_hex():
    assert svc.sha256_hex(b"abc") == (
        "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
    )
